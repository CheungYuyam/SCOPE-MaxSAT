#!/usr/bin/env python3
"""Audit and summarize the supplied eight-solver 300-second workbooks.

The files with an .xls suffix are OOXML workbooks, so they are read from
bytes with openpyxl. Solver names are assigned from the frozen package/folder
mapping. This also avoids propagating the known copied Model-cell labels in
the N2023 and NFPS wrapper outputs.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import statistics
import zipfile
from collections import Counter, defaultdict
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any, Iterable

import openpyxl


CUTOFF = 300.0
EXPECTED_TRACK_COUNTS = {
    "23-wpms-300s": 160,
    "23-pms-300s": 179,
    "24-wpms-300s": 229,
    "24-pms-300s": 216,
}
SOLVER_BY_DIRECTORY = {
    "SCOPE": "SCOPE-MaxSAT",
    "NuWLS": "NuWLS",
    "SATLike3.0": "SATLike3.0",
    "BandMaxSAT": "BandMaxSAT",
    "N2023": "NuWLS-c-2023",
    "NFPS": "NuWLS-c-FPS",
    "FPS": "SPB-MaxSAT-c-FPS",
    "Band": "SPB-MaxSAT-c-Band",
}
SOLVER_ORDER = list(SOLVER_BY_DIRECTORY.values())
TRACK_ORDER = [
    "23-wpms-300s",
    "23-pms-300s",
    "24-wpms-300s",
    "24-pms-300s",
]


def normalize_instance(value: Any) -> str:
    text = str(value).strip().replace("\\", "/").rsplit("/", 1)[-1]
    while text.lower().endswith(".wcnf"):
        text = text[:-5]
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def as_float(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "success"}


def track_from_name(name: str) -> str:
    low = name.lower()
    match = re.search(r"(?:^|_)(23|24)(?:_|$)", low)
    if match is None:
        raise ValueError(f"Cannot identify release year from workbook name: {name}")
    year = match.group(1)
    kind = "wpms" if "wpms" in low else "pms"
    return f"{year}-{kind}-300s"


def exact_sign_p(left: int, right: int) -> float:
    n = left + right
    if n == 0:
        return 1.0
    tail = sum(math.comb(n, k) for k in range(0, min(left, right) + 1)) / (2**n)
    return min(1.0, 2.0 * tail)


def holm_adjust(values: list[float]) -> list[float]:
    order = sorted(range(len(values)), key=values.__getitem__)
    adjusted = [1.0] * len(values)
    running = 0.0
    m = len(values)
    for rank, index in enumerate(order):
        running = max(running, (m - rank) * values[index])
        adjusted[index] = min(1.0, running)
    return adjusted


def median_or_none(values: Iterable[float]) -> float | None:
    materialized = list(values)
    return statistics.median(materialized) if materialized else None


def mean_or_none(values: Iterable[float]) -> float | None:
    materialized = list(values)
    return statistics.fmean(materialized) if materialized else None


def read_excluded_keys(path: Path) -> set[str]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = csv.DictReader(handle)
        return {
            row["InstanceKey"]
            for row in rows
            if row.get("DevelopmentExcluded", "").lower() == "true"
        }


def parse_workbook(
    archive_member: str,
    payload: bytes,
    excluded_keys: set[str],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    parts = PurePosixPath(archive_member).parts
    directory = parts[-2]
    model = SOLVER_BY_DIRECTORY[directory]
    dataset = track_from_name(parts[-1])
    workbook = openpyxl.load_workbook(BytesIO(payload), read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    raw_rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value).strip() if value is not None else "" for value in raw_rows[0]]
    records: list[dict[str, Any]] = []
    internal_labels: Counter[str] = Counter()

    if "Instance" in headers:
        source_format = "uniform-wrapper"
        for values in raw_rows[1:]:
            row = dict(zip(headers, values))
            if not row.get("Instance"):
                continue
            original = str(row["Instance"]).strip()
            internal = row.get("Model")
            if internal:
                internal_labels[str(internal).strip()] += 1
            feasible_field = row.get("HasVerifiedSolution", row.get("HasSolution"))
            feasible = as_bool(feasible_field)
            cost = as_float(row.get("BestCost"))
            best_time = as_float(row.get("PaperTableTimeSeconds"))
            first_time = as_float(row.get("FirstFeasibleTimeSeconds"))
    else:
        source_format = "legacy-summary"
        for values in raw_rows[1:]:
            if not values or values[0] is None:
                continue
            original = str(values[0]).strip()
            status = values[3] if len(values) > 3 else None
            feasible = str(status).strip().lower() == "success"
            cost = as_float(values[1] if len(values) > 1 else None)
            best_time = as_float(values[2] if len(values) > 2 else None)
            first_time = None

            key = normalize_instance(original)
            feasible_strict = feasible and (first_time is None or first_time <= CUTOFF)
            quality = feasible and cost is not None and best_time is not None
            records.append(
                {
                    "Dataset": dataset,
                    "Model": model,
                    "ModelDirectory": directory,
                    "InstanceKey": key,
                    "InstanceOriginal": original,
                    "FeasibleReported": feasible,
                    "FeasibleStrict300": feasible_strict,
                    "BestCost": cost,
                    "BestTimeSeconds": best_time,
                    "FirstFeasibleTimeSeconds": first_time,
                    "QualityReported": quality,
                    "QualityStrict300": quality and best_time <= CUTOFF,
                    "SourceFile": archive_member,
                    "SourceFormat": source_format,
                    "DevelopmentExcluded": key in excluded_keys,
                    "ParameterDevelopmentTrack": dataset == "23-wpms-300s",
                    "PrimaryHeldOut": dataset != "23-wpms-300s" and key not in excluded_keys,
                }
            )

        audit = {
            "file": archive_member,
            "model": model,
            "dataset": dataset,
            "sheet": sheet.title,
            "source_format": source_format,
            "raw_rows": len(raw_rows) - 1,
            "instance_rows": len(records),
            "reported_feasible": sum(row["FeasibleReported"] for row in records),
            "strict300_feasible": sum(row["FeasibleStrict300"] for row in records),
            "best_events_after_300": sum(
                row["QualityReported"] and not row["QualityStrict300"] for row in records
            ),
            "first_feasible_events_after_300": sum(
                row["FirstFeasibleTimeSeconds"] is not None
                and row["FirstFeasibleTimeSeconds"] > CUTOFF
                for row in records
            ),
            "duplicate_instance_keys": len(records) - len({row["InstanceKey"] for row in records}),
            "internal_model_labels": dict(internal_labels),
            "internal_label_matches_frozen_name": not internal_labels
            or set(internal_labels) == {model},
        }
        return records, audit

    # Uniform wrapper rows share the same post-processing block.
    # Variables are assigned inside the loop above, so re-read them here.
    records = []
    for values in raw_rows[1:]:
        row = dict(zip(headers, values))
        if not row.get("Instance"):
            continue
        original = str(row["Instance"]).strip()
        feasible = as_bool(row.get("HasVerifiedSolution", row.get("HasSolution")))
        cost = as_float(row.get("BestCost"))
        best_time = as_float(row.get("PaperTableTimeSeconds"))
        first_time = as_float(row.get("FirstFeasibleTimeSeconds"))
        key = normalize_instance(original)
        feasible_strict = feasible and (first_time is None or first_time <= CUTOFF)
        quality = feasible and cost is not None and best_time is not None
        records.append(
            {
                "Dataset": dataset,
                "Model": model,
                "ModelDirectory": directory,
                "InstanceKey": key,
                "InstanceOriginal": original,
                "FeasibleReported": feasible,
                "FeasibleStrict300": feasible_strict,
                "BestCost": cost,
                "BestTimeSeconds": best_time,
                "FirstFeasibleTimeSeconds": first_time,
                "QualityReported": quality,
                "QualityStrict300": quality and best_time <= CUTOFF,
                "SourceFile": archive_member,
                "SourceFormat": source_format,
                "DevelopmentExcluded": key in excluded_keys,
                "ParameterDevelopmentTrack": dataset == "23-wpms-300s",
                "PrimaryHeldOut": dataset != "23-wpms-300s" and key not in excluded_keys,
            }
        )

    audit = {
        "file": archive_member,
        "model": model,
        "dataset": dataset,
        "sheet": sheet.title,
        "source_format": source_format,
        "raw_rows": len(raw_rows) - 1,
        "instance_rows": len(records),
        "reported_feasible": sum(row["FeasibleReported"] for row in records),
        "strict300_feasible": sum(row["FeasibleStrict300"] for row in records),
        "best_events_after_300": sum(
            row["QualityReported"] and not row["QualityStrict300"] for row in records
        ),
        "first_feasible_events_after_300": sum(
            row["FirstFeasibleTimeSeconds"] is not None
            and row["FirstFeasibleTimeSeconds"] > CUTOFF
            for row in records
        ),
        "duplicate_instance_keys": len(records) - len({row["InstanceKey"] for row in records}),
        "internal_model_labels": dict(internal_labels),
        "internal_label_matches_frozen_name": not internal_labels
        or set(internal_labels) == {model},
    }
    return records, audit


def mark_winners(records: list[dict[str, Any]]) -> None:
    by_instance: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in records:
        by_instance[(row["Dataset"], row["InstanceKey"])].append(row)
    for rows in by_instance.values():
        for suffix, quality_key in (("Reported", "QualityReported"), ("Strict300", "QualityStrict300")):
            eligible = [row for row in rows if row[quality_key]]
            best = min((row["BestCost"] for row in eligible), default=None)
            winners = [row for row in eligible if row["BestCost"] == best]
            for row in rows:
                is_winner = row in winners
                row[f"InstanceBestCost{suffix}"] = best
                row[f"Win{suffix}"] = is_winner
                row[f"UniqueWin{suffix}"] = is_winner and len(winners) == 1
                row[f"SharedWin{suffix}"] = is_winner and len(winners) > 1


def scope_rows(records: list[dict[str, Any]], scope: str) -> list[dict[str, Any]]:
    if scope in EXPECTED_TRACK_COUNTS:
        return [row for row in records if row["Dataset"] == scope]
    if scope == "Primary held-out":
        return [row for row in records if row["PrimaryHeldOut"]]
    if scope == "All four tracks":
        return records
    raise ValueError(scope)


def solver_summary(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    scopes = TRACK_ORDER + ["Primary held-out", "All four tracks"]
    for scope in scopes:
        selected = scope_rows(records, scope)
        instances = len({(row["Dataset"], row["InstanceKey"]) for row in selected})
        for model in SOLVER_ORDER:
            rows = [row for row in selected if row["Model"] == model]
            winning_times = [row["BestTimeSeconds"] for row in rows if row["WinReported"]]
            strict_times = [row["BestTimeSeconds"] for row in rows if row["WinStrict300"]]
            feasible = sum(row["FeasibleReported"] for row in rows)
            strict_feasible = sum(row["FeasibleStrict300"] for row in rows)
            output.append(
                {
                    "Dataset": scope,
                    "Instances": instances,
                    "Model": model,
                    "Feasible": feasible,
                    "FeasiblePct": 100.0 * feasible / instances,
                    "Strict300Feasible": strict_feasible,
                    "Strict300FeasiblePct": 100.0 * strict_feasible / instances,
                    "Wins": sum(row["WinReported"] for row in rows),
                    "UniqueWins": sum(row["UniqueWinReported"] for row in rows),
                    "SharedWins": sum(row["SharedWinReported"] for row in rows),
                    "MeanBestTimeOnWinsSeconds": mean_or_none(winning_times),
                    "MedianBestTimeOnWinsSeconds": median_or_none(winning_times),
                    "Strict300Wins": sum(row["WinStrict300"] for row in rows),
                    "Strict300UniqueWins": sum(row["UniqueWinStrict300"] for row in rows),
                    "Strict300SharedWins": sum(row["SharedWinStrict300"] for row in rows),
                    "Strict300MeanBestTimeOnWinsSeconds": mean_or_none(strict_times),
                }
            )
    return output


def pairwise_summary(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    scopes = TRACK_ORDER + ["Primary held-out", "All four tracks"]
    for scope in scopes:
        selected = scope_rows(records, scope)
        indexed = {
            (row["Model"], row["Dataset"], row["InstanceKey"]): row for row in selected
        }
        instance_ids = sorted({(row["Dataset"], row["InstanceKey"]) for row in selected})
        scope_rows_out = []
        for baseline in SOLVER_ORDER[1:]:
            pairs = [
                (indexed[(SOLVER_ORDER[0], *identifier)], indexed[(baseline, *identifier)])
                for identifier in instance_ids
            ]
            scope_only = sum(a["FeasibleStrict300"] and not b["FeasibleStrict300"] for a, b in pairs)
            baseline_only = sum(b["FeasibleStrict300"] and not a["FeasibleStrict300"] for a, b in pairs)
            both = sum(a["FeasibleStrict300"] and b["FeasibleStrict300"] for a, b in pairs)
            neither = len(pairs) - scope_only - baseline_only - both
            quality_pairs = [
                (a, b) for a, b in pairs if a["QualityStrict300"] and b["QualityStrict300"]
            ]
            wins = sum(a["BestCost"] < b["BestCost"] for a, b in quality_pairs)
            losses = sum(a["BestCost"] > b["BestCost"] for a, b in quality_pairs)
            ties = len(quality_pairs) - wins - losses
            relative = []
            changed = []
            for a, b in quality_pairs:
                denominator = abs(b["BestCost"])
                if denominator == 0:
                    continue
                value = 100.0 * (b["BestCost"] - a["BestCost"]) / denominator
                relative.append(value)
                if a["BestCost"] != b["BestCost"]:
                    changed.append(value)
            row = {
                "Dataset": scope,
                "Baseline": baseline,
                "Pairs": len(pairs),
                "ScopeFeasible": sum(a["FeasibleStrict300"] for a, _ in pairs),
                "BaselineFeasible": sum(b["FeasibleStrict300"] for _, b in pairs),
                "DeltaFeasible": sum(a["FeasibleStrict300"] for a, _ in pairs)
                - sum(b["FeasibleStrict300"] for _, b in pairs),
                "BothFeasible": both,
                "ScopeOnlyFeasible": scope_only,
                "BaselineOnlyFeasible": baseline_only,
                "NeitherFeasible": neither,
                "FeasibilityExactP": exact_sign_p(scope_only, baseline_only),
                "ScopeQualityWins": wins,
                "ScopeQualityLosses": losses,
                "QualityTies": ties,
                "QualityExactP": exact_sign_p(wins, losses),
                "MedianRelativeImprovementPct": median_or_none(relative),
                "MedianRelativeImprovementChangedPct": median_or_none(changed),
                "FeasibilityHolmP": None,
                "QualityHolmP": None,
            }
            output.append(row)
            scope_rows_out.append(row)
        if scope in {"Primary held-out", "All four tracks"}:
            feasibility = holm_adjust([row["FeasibilityExactP"] for row in scope_rows_out])
            quality = holm_adjust([row["QualityExactP"] for row in scope_rows_out])
            for row, feasibility_p, quality_p in zip(scope_rows_out, feasibility, quality):
                row["FeasibilityHolmP"] = feasibility_p
                row["QualityHolmP"] = quality_p
    return output


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def build_budget_comparison(
    summary_300: list[dict[str, Any]],
    summary_60_path: Path,
) -> list[dict[str, Any]]:
    with summary_60_path.open("r", encoding="utf-8-sig", newline="") as handle:
        summary_60 = list(csv.DictReader(handle))
    lookup_60 = {(row["Dataset"], row["Model"]): row for row in summary_60}
    output = []
    for row_300 in summary_300:
        dataset_300 = row_300["Dataset"]
        if dataset_300 in EXPECTED_TRACK_COUNTS:
            dataset_60 = dataset_300.replace("300s", "60s")
        else:
            dataset_60 = dataset_300
        row_60 = lookup_60[(dataset_60, row_300["Model"])]
        output.append(
            {
                "Scope": dataset_300,
                "Model": row_300["Model"],
                "Instances": row_300["Instances"],
                "Feasible60": int(row_60["Feasible"]),
                "Feasible300": row_300["Feasible"],
                "DeltaFeasible300Minus60": row_300["Feasible"] - int(row_60["Feasible"]),
                "Wins60": int(row_60["Wins"]),
                "Wins300": row_300["Wins"],
                "DeltaWins300Minus60": row_300["Wins"] - int(row_60["Wins"]),
            }
        )
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("archive", type=Path)
    parser.add_argument("output_dir", type=Path)
    parser.add_argument("--cleaned-60", required=True, type=Path)
    parser.add_argument("--summary-60", required=True, type=Path)
    args = parser.parse_args()

    excluded_keys = read_excluded_keys(args.cleaned_60)
    if not excluded_keys:
        raise SystemExit("The exclusion set is empty; provide the normalized 60-s data used by the study")

    records: list[dict[str, Any]] = []
    file_audits = []
    with zipfile.ZipFile(args.archive) as archive:
        members = [
            name for name in archive.namelist()
            if not name.endswith("/") and PurePosixPath(name).suffix.lower() in {".xls", ".xlsx"}
        ]
        for member in sorted(members):
            workbook_records, audit = parse_workbook(
                member, archive.read(member), excluded_keys
            )
            records.extend(workbook_records)
            file_audits.append(audit)

    if len(file_audits) != 32:
        raise SystemExit(f"Expected 32 workbooks, found {len(file_audits)}")
    expected_records = sum(EXPECTED_TRACK_COUNTS.values()) * len(SOLVER_ORDER)
    if len(records) != expected_records:
        raise SystemExit(f"Expected {expected_records} records, found {len(records)}")

    coverage_errors = []
    for dataset, expected in EXPECTED_TRACK_COUNTS.items():
        key_sets = {
            model: {row["InstanceKey"] for row in records if row["Dataset"] == dataset and row["Model"] == model}
            for model in SOLVER_ORDER
        }
        for model, keys in key_sets.items():
            if len(keys) != expected:
                coverage_errors.append(f"{dataset}/{model}: {len(keys)} != {expected}")
        reference = key_sets[SOLVER_ORDER[0]]
        for model, keys in key_sets.items():
            if keys != reference:
                coverage_errors.append(f"{dataset}/{model}: instance set differs from SCOPE")
    if coverage_errors:
        raise SystemExit("; ".join(coverage_errors))

    mark_winners(records)
    summary = solver_summary(records)
    pairwise = pairwise_summary(records)
    budget = build_budget_comparison(summary, args.summary_60)

    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.output_dir / "all_300s_cleaned_results_long.csv", records)
    write_csv(args.output_dir / "all_300s_solver_summary.csv", summary)
    write_csv(args.output_dir / "all_300s_pairwise_vs_scope.csv", pairwise)
    write_csv(args.output_dir / "all_60s_vs_300s_summary.csv", budget)

    audit = {
        "archive": str(args.archive),
        "archive_sha256": hashlib.sha256(args.archive.read_bytes()).hexdigest(),
        "cutoff_seconds": CUTOFF,
        "models": SOLVER_ORDER,
        "datasets": EXPECTED_TRACK_COUNTS,
        "excluded_key_count": len(excluded_keys),
        "workbook_count": len(file_audits),
        "record_count": len(records),
        "coverage_errors": coverage_errors,
        "best_events_after_300": sum(item["best_events_after_300"] for item in file_audits),
        "first_feasible_events_after_300": sum(
            item["first_feasible_events_after_300"] for item in file_audits
        ),
        "internal_label_mismatch_files": [
            item["file"] for item in file_audits
            if not item["internal_label_matches_frozen_name"]
        ],
        "files": file_audits,
    }
    with (args.output_dir / "all_300s_data_audit.json").open("w", encoding="utf-8") as handle:
        json.dump(audit, handle, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    main()
